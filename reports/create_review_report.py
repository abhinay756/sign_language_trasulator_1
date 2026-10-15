from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
ASSET_DIR = REPORT_DIR / "assets"
OUTPUT = REPORT_DIR / "HandSign_Review_Report.docx"
metrics = json.loads((REPORT_DIR / "metrics.json").read_text(encoding="utf-8"))
classes = metrics["classes"]
counts = metrics["sample_counts"]
results = metrics["results"]

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F4F7FA"
GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRAY, size="6") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_font(run, name="Aptos", size=10.5, bold=False, color=BLACK, italic=False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_text(cell, text, bold=False, color=BLACK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=RGBColor(255, 255, 255), size=9)
        set_cell_shading(table.rows[0].cells[i], NAVY)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8.8)
            if row_index % 2 == 1:
                set_cell_shading(cells[i], PALE)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_font(run, size=10)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    set_font(run, size=9, italic=True, color=RGBColor(70, 70, 70))


def add_picture(doc, path: Path, width: float, caption: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=8, color=RGBColor(90, 90, 90))
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Title" else "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        ppr = style._element.pPr
        if ppr is not None:
            paragraph_borders = ppr.find(qn("w:pBdr"))
            if paragraph_borders is not None:
                ppr.remove(paragraph_borders)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_title(doc):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("Real Time Hand Sign Recognition Project Report")
    set_font(run, name="Aptos Display", size=25, bold=True)
    ppr = paragraph._p.get_or_add_pPr()
    paragraph_borders = ppr.find(qn("w:pBdr"))
    if paragraph_borders is not None:
        ppr.remove(paragraph_borders)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Review aligned report for classroom evaluation")
    set_font(run, size=12, color=RGBColor(60, 60, 60))
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    rows = [
        ("Project", "HandSign: real-time sign and alphabet recognition"),
        ("Team", "Add 4 to 6 members from the given namelist"),
        ("Technology", "MediaPipe Hands, PyTorch, OpenCV, classical and quantum simulation baselines"),
        ("Classes", ", ".join(classes)),
    ]
    for r, (key, value) in enumerate(rows):
        set_cell_text(table.cell(r, 0), key, bold=True, size=9.5)
        set_cell_shading(table.cell(r, 0), BLUE)
        set_cell_text(table.cell(r, 1), value, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    add_body(doc, "This report explains the problem, literature basis, data pipeline, model choices, experiments, results, and comparison plan for the HandSign project. It is structured so the same document can be used as the Review 1 introduction and literature section, the Review 2 model execution record, and the Review 3 comparison and rough report draft.")


def add_review_schedule(doc):
    doc.add_heading("Review Schedule and Deliverables", level=1)
    add_table(doc, ["Review", "Date and mode", "Required evidence", "Where covered"], [
        ("Review 1", "20 marks\nOffline classroom\n10 August 2026", "Team formation, introduction, literature review", "Sections 1 and 2"),
        ("Review 2", "30 marks\nOffline classroom\n25 August 2026", "ANN, QNN, CNN execution; code explanation; accuracy, loss graphs, confusion matrix", "Sections 3 to 6"),
        ("Review 3", "30 marks\nOffline classroom\n11 September 2026", "Two-model comparison with images, accuracy, loss graphs, confusion matrices, rough report", "Sections 7 and 8"),
    ], widths=[0.9, 1.4, 2.6, 1.6])


def add_review_one(doc):
    doc.add_heading("1 Review 1 Team Formation and Introduction", level=1)
    doc.add_heading("1.1 Team formation", level=2)
    add_body(doc, "The team should contain a minimum of four and a maximum of six students selected from the supplied namelist. The table below is intentionally left editable so the final submission can contain the approved names, roll numbers, and responsibilities.")
    add_table(doc, ["Member", "Name", "Roll number", "Responsibility"], [
        ("1", "", "", "Project coordination and presentation"),
        ("2", "", "", "Data collection and annotation"),
        ("3", "", "", "Model development and training"),
        ("4", "", "", "Evaluation and documentation"),
        ("5", "", "", "Optional: quantum model and literature"),
        ("6", "", "", "Optional: interface and testing"),
    ], widths=[0.55, 2.2, 1.1, 2.65])
    doc.add_heading("1.2 Problem statement", level=2)
    add_body(doc, "People who communicate with sign language may need a fast, low-cost interface that converts a small vocabulary of hand signs into readable labels. The proposed system uses a normal webcam to recognize alphabet signs and common interaction gestures without requiring a glove, depth sensor, or cloud service.")
    doc.add_heading("1.3 Aim and objectives", level=2)
    add_bullet(doc, "Build a real-time webcam system for eight labels: A, B, C, HI, BYE, THANK_YOU, YES, and NO.")
    add_bullet(doc, "Normalize hand landmarks so the model is less sensitive to translation and scale in the camera frame.")
    add_bullet(doc, "Compare a temporal GRU project model with ANN, CNN, and a state-vector QNN simulation.")
    add_bullet(doc, "Report accuracy, loss behavior, and class-level errors using confusion matrices.")
    add_bullet(doc, "Provide a reproducible training and inference pipeline that can be extended to more alphabet letters.")
    doc.add_heading("1.4 Scope and users", level=2)
    add_body(doc, "The first version is a classroom research prototype, not a complete sign-language translator. It recognizes a fixed vocabulary and assumes one visible hand. It can be extended to all alphabet letters, two-hand signs, sentence decoding, and language-specific sign standards after collecting a larger and more diverse dataset.")


def add_literature(doc):
    doc.add_heading("2 Literature Review", level=1)
    add_body(doc, "The literature supports a staged design: a real-time hand tracker reduces the raw video to a compact representation; temporal models capture motion; convolutional models provide a strong spatial baseline; encoder-decoder CNNs are appropriate when the task includes pixel-level hand segmentation; and variational quantum circuits provide a research extension for low-dimensional features.")
    add_table(doc, ["Area", "Representative work", "Contribution to this project"], [
        ("Hand tracking", "Zhang et al., MediaPipe Hands, 2020", "On-device hand landmark detection makes webcam inference lightweight and gives 21 three-dimensional landmarks per hand."),
        ("Image segmentation", "Ronneberger et al., U-Net, 2015", "Provides the encoder-decoder design to segment the hand region when pixel masks are available."),
        ("Attention segmentation", "Oktay et al., Attention U-Net, 2018", "Shows how attention gates can suppress irrelevant background regions."),
        ("Nested segmentation", "Zhou et al., UNet++, 2018", "Uses nested skip pathways to reduce the semantic gap between encoder and decoder features."),
        ("Temporal recognition", "GRU and recurrent sequence modeling", "Motivates learning a sequence representation for HI, BYE, and THANK_YOU motion gestures."),
        ("Quantum learning", "Variational quantum circuit literature", "Motivates a small state-vector QNN simulation after PCA reduces the feature vector to four angles."),
    ], widths=[1.25, 2.35, 3.0])
    doc.add_heading("2.1 References for presentation", level=2)
    references = [
        "Zhang, F. et al. MediaPipe Hands: On-device Real-time Hand Tracking. arXiv:2006.10214. https://arxiv.org/abs/2006.10214",
        "Ronneberger, O., Fischer, P., and Brox, T. U-Net: Convolutional Networks for Biomedical Image Segmentation. arXiv:1505.04597. https://arxiv.org/abs/1505.04597",
        "Oktay, O. et al. Attention U-Net: Learning Where to Look for the Pancreas. arXiv:1804.03999. https://arxiv.org/abs/1804.03999",
        "Zhou, Z. et al. UNet++: A Nested U-Net Architecture for Medical Image Segmentation. arXiv:1807.10165. https://arxiv.org/abs/1807.10165",
        "Schuld, M. and Petruccione, F. Supervised Learning with Quantum Computers. Springer, 2018.",
    ]
    for reference in references:
        add_bullet(doc, reference)


def add_system(doc):
    doc.add_heading("3 Proposed System and Dataset", level=1)
    add_picture(doc, ASSET_DIR / "pipeline.png", 6.7, "Figure 1. HandSign pipeline from webcam frames to a stable class prediction.")
    add_body(doc, "The webcam frame is processed by MediaPipe Hands. The first detected hand is converted into 63 normalized values: 21 landmarks multiplied by x, y, and z coordinates. The wrist is used as the origin and the largest 2D landmark radius is used for scale normalization. Thirty consecutive frames form one training example. This representation makes the same sign less dependent on where the hand appears in the frame.")
    add_table(doc, ["Dataset item", "Value"], [
        ("Classes", ", ".join(classes)),
        ("Total sequences", str(sum(counts.values()))),
        ("Sequence length", "30 frames"),
        ("Features per frame", "63 normalized landmark values"),
        ("Validation split", "20 percent, stratified, random seed 7"),
        ("Runtime", "CPU webcam inference with OpenCV"),
    ], widths=[2.2, 4.4])
    add_body(doc, "The collected data contains " + ", ".join(f"{label}: {counts[label]}" for label in classes) + ". The split is stratified so every label is represented in both training and validation sets.")
    doc.add_heading("3.1 Reproduction commands", level=2)
    for command in [
        "uv run collect-signs --samples 40",
        "uv run train-signs --epochs 40",
        "uv run evaluate-signs --checkpoint artifacts/gesture_model.pt",
        "uv run realtime-signs --checkpoint artifacts/gesture_model.pt",
    ]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(command)
        set_font(run, name="Consolas", size=9.5, color=RGBColor(30, 70, 120))


def add_review_two(doc):
    doc.add_heading("4 Review 2 Model Execution", level=1)
    add_body(doc, "Four models were evaluated on the same stratified validation split. The GRU is the project model used for real-time recognition. The ANN and CNN are supervised baselines. The QNN is a differentiable four-qubit state-vector simulation, which is suitable for offline classroom execution without quantum hardware.")
    add_table(doc, ["Model", "Input representation", "Main layers", "Purpose"], [
        ("ANN", "Flattened 30 x 63 sequence", "Dense 128, Dense 64, softmax", "Fully connected baseline"),
        ("QNN simulation", "Four PCA angles", "RY rotations, ring CNOTs, Z measurements, linear head", "Quantum learning comparison"),
        ("CNN", "3 x 30 x 21 landmark tensor", "Two convolution blocks, pooling, adaptive average pooling", "Spatial feature baseline"),
        ("GRU project model", "30 x 63 sequence", "Projection, two-layer bidirectional GRU, classifier", "Final real-time model"),
    ], widths=[1.25, 1.8, 2.5, 1.35])
    doc.add_heading("4.1 ANN code explanation", level=2)
    add_body(doc, "The ANN receives the complete landmark sequence after flattening. Dense layers with ReLU learn nonlinear combinations of finger and temporal values, while the final softmax produces one probability per class. It is easy to train, but flattening removes the explicit time structure of motion gestures.")
    doc.add_heading("4.2 QNN code explanation", level=2)
    add_body(doc, "The QNN first applies PCA to reduce the 1,890-value sequence to four values. These become rotation angles for four simulated qubits. Each circuit layer applies trainable RY rotations followed by a ring of controlled-NOT gates. Z-axis expectation values are measured and passed to a small linear classifier. This is a software simulation of a variational quantum circuit; no quantum hardware is claimed.")
    doc.add_heading("4.3 CNN code explanation", level=2)
    add_body(doc, "The CNN reshapes the landmark sequence into three channels with a 30 by 21 spatial grid. Convolution layers learn local patterns across time and landmarks, pooling reduces the representation, and global average pooling feeds the classifier. This is a classification CNN, not a U-Net segmentation network.")
    doc.add_heading("4.4 U-Net requirement and technically correct extension", level=2)
    add_body(doc, "U-Net, Attention U-Net, U-Net++, and ResUNet are segmentation architectures. They require pixel-level masks, whereas the current HandSign dataset stores landmark sequences. The recommended extension is to collect hand masks or generate reviewed masks from webcam frames, train an Attention U-Net to segment the hand ROI, and pass the segmented ROI to the classifier. The present document reports the executed landmark CNN honestly and records the U-Net stage as the next experiment rather than claiming an unexecuted segmentation result.")


def add_results(doc):
    doc.add_heading("5 Results Accuracy Loss and Confusion Matrices", level=1)
    add_table(doc, ["Model", "Validation accuracy", "Interpretation"], [
        ("GRU project model", f"{results['GRU project model'] * 100:.1f}%", "Best overall result; preserves temporal structure for moving signs."),
        ("ANN", f"{results['ANN'] * 100:.1f}%", "Useful baseline; loses explicit temporal relationships."),
        ("CNN", f"{results['CNN'] * 100:.1f}%", "Lower result on this landmark tensor; needs stronger spatial encoding or image input."),
        ("QNN simulation", f"{results['QNN simulation'] * 100:.1f}%", "Research baseline; four-angle compression is highly restrictive."),
    ], widths=[1.65, 1.45, 3.8])
    add_picture(doc, ASSET_DIR / "accuracy_comparison.png", 6.45, "Figure 2. Accuracy comparison for the executed models.")
    add_picture(doc, ASSET_DIR / "training_curves.png", 6.55, "Figure 3. Training loss and validation accuracy curves for ANN, CNN, and QNN simulation.")
    add_picture(doc, ASSET_DIR / "confusion_matrices.png", 6.6, "Figure 4. Confusion matrices on the common validation split.")
    add_body(doc, "The GRU is selected for deployment because it achieved the highest validation accuracy and models the temporal patterns of HI, BYE, and THANK_YOU. The ANN remains valuable as a simple baseline. The CNN and QNN results are useful research evidence: they show that a model can be weaker when the representation does not match its inductive bias or when dimensionality reduction removes too much information.")


def add_review_three(doc):
    doc.add_heading("6 Review 3 Model Comparison and Rough Report", level=1)
    add_body(doc, "The primary comparison is between the deployed GRU and the ANN baseline because both use the same landmark sequence input. This makes the comparison about model structure rather than a different sensor or dataset.")
    add_table(doc, ["Comparison item", "GRU project model", "ANN baseline"], [
        ("Input", "30 ordered frames x 63 features", "Same sequence flattened into one vector"),
        ("Temporal information", "Explicitly learned through bidirectional GRU states", "Implicit only through fixed feature positions"),
        ("Validation accuracy", f"{results['GRU project model'] * 100:.1f}%", f"{results['ANN'] * 100:.1f}%"),
        ("Best use", "Real-time sequence recognition", "Simple benchmark and ablation study"),
        ("Main limitation", "Needs enough diverse sequences and careful smoothing", "Motion order is not modeled naturally"),
    ], widths=[1.7, 2.35, 2.85])
    add_picture(doc, ASSET_DIR / "pipeline.png", 6.55, "Figure 5. System image for the final project presentation.")
    add_body(doc, "The confusion matrices should be discussed class by class during the viva. Diagonal values indicate correct classifications. Off-diagonal values identify signs that look similar under the current camera and hand pose. The next data collection round should focus on the most frequently confused pair, vary the signer and lighting, and include a neutral no-hand sequence if false positives are observed.")
    doc.add_heading("6.1 Rough draft of the final report", level=2)
    add_table(doc, ["Chapter", "Content to submit"], [
        ("1. Introduction", "Problem, motivation, objectives, scope, and contribution."),
        ("2. Literature review", "MediaPipe hand tracking, temporal models, CNN/U-Net segmentation, and quantum learning."),
        ("3. Methodology", "Data collection, landmark normalization, sequence creation, model architectures, and training settings."),
        ("4. Implementation", "Folder structure, commands, key code blocks, webcam interface, and checkpoint format."),
        ("5. Results", "Accuracy table, loss curves, confusion matrices, sample images, and model comparison."),
        ("6. Discussion", "Why GRU performs best, errors, limitations, ethical considerations, and U-Net extension."),
        ("7. Conclusion", "Summary of the working prototype and next steps toward a larger vocabulary."),
    ], widths=[1.7, 4.9])
    doc.add_heading("6.2 Presentation sequence", level=2)
    for item in [
        "Show the problem and the eight target labels.",
        "Explain how a webcam frame becomes 63 normalized landmark values.",
        "Show the pipeline image and a short live demonstration.",
        "Explain ANN, QNN simulation, CNN, and GRU with one slide each.",
        "Present accuracy, loss curves, and confusion matrices.",
        "Close with limitations, U-Net segmentation extension, and future alphabet expansion.",
    ]:
        add_bullet(doc, item)


def add_conclusion(doc):
    doc.add_heading("7 Conclusion and Next Work", level=1)
    add_body(doc, "HandSign is a complete classroom-ready prototype for real-time recognition of a small sign vocabulary. The current evidence shows that a temporal GRU is the strongest of the tested models for this landmark-sequence task. The project also includes an ANN baseline, a state-vector QNN simulation, and a CNN baseline so the team can explain model selection with measured results rather than only architecture diagrams.")
    add_body(doc, "Before final submission, replace the team-member blanks, add screenshots from the live webcam application, document the exact hardware and software versions, and complete the U-Net segmentation extension if the evaluator interprets the CNN requirement as specifically requiring a U-Net-family model.")


def main():
    doc = Document()
    style_document(doc)
    add_title(doc)
    add_review_schedule(doc)
    add_review_one(doc)
    add_literature(doc)
    add_system(doc)
    add_review_two(doc)
    add_results(doc)
    add_review_three(doc)
    add_conclusion(doc)
    doc.core_properties.title = "Real Time Hand Sign Recognition Project Report"
    doc.core_properties.subject = "Review 1 Review 2 and Review 3 project report"
    doc.core_properties.author = "HandSign Project Team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
